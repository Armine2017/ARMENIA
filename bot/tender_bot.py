#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
ՏԵՆԴԵՐԱՅԻՆ ՄՈՆԻՏՈՐ v2 — ամենօրյա ավտոմատ բեռնիչ
=====================================================
Ինչ է անում.
  * Ստուգում է config.json-ում նշված կայքերը (ներառյալ ARMEPS «Հայտեր»-ը և
    Ֆիննախի «Մրցույթի հայտարարություն և հրավեր»-ը),
  * կարդում է էջերի շրջանցումը (էջ 1, 2, 3 …) և կանգնում, երբ հասնում է since_date-ին,
  * պահպանում է միայն ՆՈՐ (և հոկտեմբերից սկսած) հայտարարությունները՝
    data/<ԱՄՍԱԹԻՎ>/<ԿԱՏԵԳՈՐԻԱ>/... թղթապանակներում,
  * դուրս է գրում ապրանքը/առարկան, կազմակերպությունը, ծածկագիրը, վերջնաժամկետը,
    ՀԵՌԱԽՈՍԸ և ԷԼ. ՓՈՍՏԸ reports/*.csv-ում,
  * առանձին աղյուսակ է գրում «Հայտեր»-ը (մասնակից, ՀՎՀՀ, երկիր, հայտի գումար),
  * վերցնում է նաև inbox/ թղթապանակը. եթե զննարկիչով պահել ես .html/.xlsx էջ,
    ծրագիրը ինքը կկարդա այն (JS-ով բեռնվող էջերի համար)։

Գործարկում.
  python tender_bot.py                 # մեկ անգամ ստուգել բոլոր կայքերը
  python tender_bot.py --demo          # փորձնական՝ լոկալ օրինակներով (առանց ինտերնետի)
  python tender_bot.py --archive       # հոկտեմբերից սկսած ամբողջ արխիվը (բոլոր էջերով)
  python tender_bot.py --since 2026-10-01
  python tender_bot.py --loop          # ամեն օր config-ի run_time-ին ինքն է աշխատում
  python tender_bot.py --reset         # մոռանալ հին բեռնումները
"""

import argparse
import csv
import datetime as dt
import hashlib
import html as htmllib
import json
import logging
import re
import sys
import time
import urllib.parse
import urllib.request
from pathlib import Path

BASE = Path(__file__).resolve().parent
CONFIG_PATH = BASE / "config.json"
SEEN_PATH = BASE / "seen.json"
LOG_DIR = BASE / "logs"
UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) TenderMonitor/2.0"

# ------------------------------------------------------------------ կարգավորումներ
DEFAULT_CONFIG = {
    "since_date": "2026-10-01",          # բեռնել միայն այս օրվանից սկսած հայտարարությունները
    "run_time": "08:00",                 # ամենօրյա գործարկման ժամը (--loop ռեժիմ)
    "delay_seconds": 2,                  # դադար հարցումների միջև
    "fetch_detail_pages": True,          # բացել նաև հայտարարության էջը (հեռախոս/էլ.փոստ)
    "only_keywords": [],                 # օր. ["ճանապարհ"] — դատարկ = բոլորը
    "max_pages": 25,                     # մինչև քանի էջ շրջանցի (1-ական՝ ամեն օր)
    "inbox_folder": "inbox",
    "categories": {
        "Շինարարություն": ["շին", "կառուց", "նորոգ", "ասֆալտ", "շենք", "կամուրջ", "ջրագիծ",
                            "պուրակ", "զբոսայգի", "տանիք", "քանդ"],
        "ՏՏ և ծրագրակազմ": ["համակարգչ", "ծրագր", "սերվեր", "տեղեկատվական", "ինտերնետ", "ցանց"],
        "Բժշկական": ["բժշկ", "դեղ", "հիվանդանոց", "բուժ", "կլինիկ", "լաբորատոր"],
        "Տրանսպորտ": ["տրանսպորտ", "մեքենա", "ավտոմեքենա", "ճանապարհ", "վառելիք", "ավտոբուս",
                      "բենզին", "շարժասանդուղք", "վերելակ", "մետրո"],
        "Ծառայություններ": ["ծառայություն", "սպասարկում", "հսկող", "խորհրդատվ", "մաքր"],
        "Ապրանքներ": ["մատակարարում", "ձեռքբերում", "գնում", "ապրանք", "սարքավորում", "գույք"],
        "Այլ": [],
    },
    "sources": [
        {
            "name": "ARMEPS-Հայտեր",
            "url": "https://armeps.am/ppcm/public/bid-report",
            "kind": "js",                                    # JS-ով բեռնվող աղյուսակ
            "page_pattern": "https://armeps.am/ppcm/public/bid-report?page={n}",
            "parse": "table",
            "date_order": "MDY",                             # 08/21/2026 = օգոստոսի 21
        },
        {
            "name": "MinFin-Մրցույթի-հայտարարություն",
            "url": "https://gnumner.minfin.am/hy/page/bac_mrcuyti_haytararutyun_ev_hraver/",
            "kind": "html",
            "page_pattern": "https://gnumner.minfin.am/hy/page/bac_mrcuyti_haytararutyun_ev_hraver/{n}",
            "parse": "tender-ad",
        },
        {
            "name": "MinFin-Հայտարարություններ",
            "url": "https://gnumner.minfin.am/hy/page/gnumneri_haytararutyunner_/",
            "kind": "html",
            "page_pattern": "https://gnumner.minfin.am/hy/page/gnumneri_haytararutyunner_/{n}",
            "parse": "tender-ad",
        },
    ],
}

TENDER_COLUMNS = [
    "Ամսաթիվ", "Աղբյուր", "Կատեգորիա", "Ապրանք / Առարկա (նպատակը)", "Կազմակերպություն",
    "Ծածկագիր", "Վերջնաժամկետ", "Կարգավիճակ", "Հեռախոս", "Էլ. փոստ", "Հղում", "Պահպանված ֆայլ",
]
BID_COLUMNS = ["Ամսաթիվ", "Աղբյուր", "Մրցույթի ծածկագիր", "Մրցույթի անվանում",
               "Մասնակից", "ՀՎՀՀ / Անձնագիր", "Երկիր", "Հայտ"]

# ------------------------------------------------- տվյալների դուրսգրման կանոններ
PHONE_RE = re.compile(
    r"(?:\+374[\s.\-]?\d{2}[\s.\-]?\d{3}[\s.\-]?\d{3}"
    r"|\(\+374[\s.\-]?\d{2}\)[\s.\-]?\d{3}[\s.\-]?\d{3}"
    r"|0\d{2}[\s.\-]?\d{3}[\s.\-]?\d{3})"
)
EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")
DATE_RE = re.compile(r"\d{1,2}[./]\d{1,2}[./]\d{2,4}|\d{4}-\d{2}-\d{2}")
ORG_RE = re.compile(r"(?:Պատվիրատու|պատվիրատու|Կազմակերպություն|կազմակերպություն)\s*[:\-–՝]?\s*([^,|\n]{3,80})")
DEADLINE_RE = re.compile(r"(?:ժամկետ|մինչև|մինչեւ|deadline)\D{0,40}?(\d{1,2}[./\-]\d{1,2}[./\-]\d{2,4})", re.I)
CODE_RE = re.compile(r"[A-ZԱ-Ֆ0-9]{2,}[‐\-–/][A-ZԱ-Ֆ0-9/‐\-–]{4,}")
# «(Հրապարակված է 2026-10-01 13:30:00-ից մինչև 2026-11-01 11:00:00 ժամը ներառյալ)»
AD_PERIOD_RE = re.compile(
    r"Հրապարակված\s+է\s*(\d{4}-\d{2}-\d{2})(?:[ T]\d{2}:\d{2}(?::\d{2})?)?\s*[-–]ից\s*"
    r"մինչ[ևե]ւ\s*(\d{4}-\d{2}-\d{2})(?:[ T]\d{2}:\d{2}(?::\d{2})?)?"
)
FIELD_HINTS = {
    "org": ["պատվիրատու", "կազմակերպ", "հաճախորդ", "buyer", "authname"],
    "code": ["ծածկագիր", "code", "ծածկագր"],
    "subject": ["անվանում", "առարկ", "ապրանք", "նպատակ", "վերնագիր", "նկարագր", "description", "title"],
    "date": ["հրատարակ", "հրապարակ", "ամսաթիվ", "date", "տեղադրման"],
    "deadline": ["վերջնաժամկետ", "ժամկետ", "deadline", "open"],
    "status": ["կարգավիճակ", "status"],
    "phone": ["հեռախոս", "phone", "tel"],
    "email": ["էլ", "փոստ", "mail", "@"],
    "url": ["հղում", "link", "url"],
    "subject_code": ["մրցույթի ծածկագիր"],
}

log = logging.getLogger("tender_bot")


# ------------------------------------------------------------------ օժանդակ
def setup_logging():
    LOG_DIR.mkdir(exist_ok=True)
    fmt = logging.Formatter("%(asctime)s | %(levelname)-7s | %(message)s", "%H:%M:%S")
    fh = logging.FileHandler(LOG_DIR / f"run-{dt.datetime.now():%Y-%m-%d_%H%M%S}.log", encoding="utf-8")
    fh.setFormatter(fmt)
    sh = logging.StreamHandler(sys.stdout)
    sh.setFormatter(fmt)
    log.setLevel(logging.INFO)
    log.handlers.clear()
    log.addHandler(fh)
    log.addHandler(sh)


def load_config():
    if not CONFIG_PATH.exists():
        CONFIG_PATH.write_text(json.dumps(DEFAULT_CONFIG, ensure_ascii=False, indent=2), encoding="utf-8")
        log.info("Ստեղծվեց նոր config.json — կարող եք խմբագրել Notepad-ով։")
    cfg = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
    for key, val in DEFAULT_CONFIG.items():
        cfg.setdefault(key, val)
    return cfg


def load_seen():
    return json.loads(SEEN_PATH.read_text(encoding="utf-8")) if SEEN_PATH.exists() else {}


def sanitize(name, limit=60):
    name = re.sub(r'[\\/:*?"<>|\r\n\t]+', " ", str(name))
    name = re.sub(r"\s+", " ", name).strip(" .")
    return name[:limit] or "անանուն"


def strip_tags(fragment):
    fragment = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", fragment, flags=re.S | re.I)
    fragment = re.sub(r"<br\s*/?>|</(p|div|li|tr|h\d)>", " ", fragment, flags=re.I)
    text = re.sub(r"<[^>]+>", " ", fragment)
    text = htmllib.unescape(text).replace("\u00a0", " ")
    return re.sub(r"\s+", " ", text).strip()


def html_with_links_marked(page_html):
    """<a href=u>Title</a> → [[u|Title]] . Վերադարձնում է հարթ տեքստ՝ հղումներով։"""
    def repl(m):
        href, inner = m.group(1), strip_tags(m.group(2))
        if not inner:
            return " "
        return f" [[{href}|{inner}]] "
    text = re.sub(r"<a\b[^>]*href=[\"']([^\"']+)[\"'][^>]*>(.*?)</a>", repl, page_html, flags=re.S | re.I)
    return strip_tags(text)


def http_get(url, timeout=40):
    req = urllib.request.Request(url, headers={"User-Agent": UA, "Accept-Language": "hy,en;q=0.8"})
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        raw = resp.read()
        enc = resp.headers.get_content_charset() if resp.headers else None
    return raw.decode(enc or "utf-8", errors="replace")


def http_get_bytes(url, timeout=60):
    req = urllib.request.Request(url, headers={"User-Agent": UA})
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        return resp.read()


def fetch(url, retries=3):
    for attempt in range(1, retries + 1):
        try:
            return http_get(url)
        except Exception as exc:
            log.warning("Չհաջողվեց բացել %s (փորձ %d/%d): %s", url, attempt, retries, exc)
            time.sleep(10 if attempt < retries else 0)
    return None


def fetch_js(url, timeout=45000):
    """JS-ով բեռնվող էջերի բացում Playwright-ով (եթե տեղադրված է)։"""
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        log.warning("Playwright տեղադրված չէ. JS-ով բեռնվող էջը չի բացվի։ Տեղադրելու համար՝ "
                    "pip install playwright  &&  playwright install chromium")
        return None
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True)
            page = browser.new_page(user_agent=UA)
            page.goto(url, wait_until="networkidle", timeout=timeout)
            page.wait_for_timeout(1500)
            content = page.content()
            browser.close()
        return content
    except Exception as exc:
        log.warning("Playwright-ը չկարողացավ բացել %s: %s", url, exc)
        return None


# ------------------------------------------------------------------ ամսաթվեր
def parse_date_any(value, order="DMY"):
    """Տարբեր ձևաչափեր → datetime.date (չհաջողվելիս՝ None)։"""
    if not value:
        return None
    s = str(value).strip()
    m = re.search(r"(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    m = re.search(r"(\d{1,2})[./](\d{1,2})[./](\d{2,4})", s)
    if m:
        a, b, year = int(m.group(1)), int(m.group(2)), int(m.group(3))
        year += 2000 if year < 100 else 0
        first, second = (a, b) if order.upper() == "DMY" else (b, a)
        try:
            return dt.date(year, second, first)
        except ValueError:
            return None
    return None


def human_date(d):
    return d.strftime("%d.%m.%Y") if isinstance(d, dt.date) else (str(d) if d else "")


def extract_fields(text):
    phones = sorted(set(PHONE_RE.findall(text or "")))
    emails = sorted(set(m.lower() for m in EMAIL_RE.findall(text or "")))
    org_m = ORG_RE.search(text or "")
    dead_m = DEADLINE_RE.search(text or "")
    code = ""
    for candidate in CODE_RE.findall(text or ""):
        if re.search(r"[A-Za-zԱ-Ֆա-ֆ]", candidate) and not DATE_RE.fullmatch(candidate.strip()):
            code = candidate
            break
    return {
        "phones": "; ".join(phones[:6]),
        "emails": "; ".join(emails[:6]),
        "org": org_m.group(1).strip() if org_m else "",
        "deadline": dead_m.group(1) if dead_m else "",
        "code": code,
        "date": DATE_RE.search(text or "").group(0) if DATE_RE.search(text or "") else "",
    }


def categorize(cfg, title, rest=""):
    """Կատեգորիան՝ կշռված. վերնագրում գտնված բառը 3 անգամ արժեքավոր է,
    քան պատվիրատուի անվան կամ այլ դաշտերի մեջ պատահաբար հանդիպած բառը։"""
    tl, rl = (title or "").lower(), (rest or "").lower()
    best, best_score = "Այլ", 0
    for cat, words in cfg["categories"].items():
        score = 0
        for w in words:
            if not w:
                continue
            wl = w.lower()
            if wl in tl:
                score += 3
            elif wl in rl:
                score += 1
        if score > best_score:
            best, best_score = cat, score
    return best


# ------------------------------------------------------------------ վերլուծիչներ
def parse_minfin_ads(page_html, base_url):
    """gnumner.minfin.am-ի «Մրցույթի հայտարարություն և հրավեր» էջերի վերլուծություն։
    Ամեն գրառում ունի հղում → armeps.am/epps/cft/listContractDocuments.do?resourceId=…,
    իսկ ներքևում՝ «(Հրապարակված է 2026-10-01 13:30:00-ից մինչև 2026-11-01 11:00:00 …)»։"""
    text = html_with_links_marked(page_html)
    items, seen_ids = [], set()
    for m in re.finditer(r"\[\[([^\]|]+)\|([^\]]+)\]\]", text):
        url = m.group(1).strip()
        title = m.group(2).strip()
        if len(title) < 15 or "armeps.am" not in url:
            continue
        window = text[m.end(): m.end() + 600].split("[[")[0]   # միայն տվյալ գրառման տեքստը
        dm = AD_PERIOD_RE.search(window)
        rid = re.search(r"resourceId=(\d+)", url)
        key = rid.group(1) if rid else title[:80]
        if key in seen_ids:
            continue
        seen_ids.add(key)
        items.append({
            "title": title,
            "url": url,
            "snippet": window[:400],
            "date": dm.group(1) if dm else "",
            "deadline": dm.group(2) if dm else "",
        })
    return items


def parse_html_table(page_html, base_url, date_order="DMY"):
    """HTML աղյուսակները → գրառումների ցանկ (ARMEPS PPCM «Հայտեր» էջի համար)։"""
    items = []
    for table in re.findall(r"<table\b.*?</table>", page_html, flags=re.S | re.I):
        headers = strip_tags(" ".join(re.findall(r"<th\b[^>]*>(.*?)</th>", table, flags=re.S | re.I)))
        header_list = [strip_tags(x) for x in re.findall(r"<th\b[^>]*>(.*?)</th>", table, flags=re.S | re.I)]
        for tr in re.findall(r"<tr\b[^>]*>(.*?)</tr>", table, flags=re.S | re.I):
            cells = [strip_tags(x) for x in re.findall(r"<td\b[^>]*>(.*?)</td>", tr, flags=re.S | re.I)]
            if len(cells) < 3:
                continue
            if all(c.replace(" ", "").isdigit() for c in cells):      # NN, 1, 2… աղյուսակ
                continue
            if cells == header_list:                                  # վերնագրային տող
                continue
            row = {}
            if header_list and len(header_list) == len(cells):
                row = dict(zip(header_list, cells))
            context = " ".join(cells)
            if any(h in headers for h in ("Մասնակից", "ՀՎՀՀ")) or "Մասնակից" in context:
                items.append({"kind": "bid", "cells": cells, "row": row, "text": context})
            else:
                items.append({"kind": "tender", "cells": cells, "row": row, "text": context,
                              "title": context[:300]})
    return items


def map_row_fields(row, date_order="DMY"):
    """Աղյուսակի սյունակները → մեր դաշտերը (ըստ վերնագրերի բառերի)։"""
    out = {"org": "", "code": "", "subject": "", "date": "", "deadline": "", "status": "",
           "phones": "", "emails": "", "url": "", "cells": []}
    for header, value in row.items():
        h = header.lower()
        for field, hints in FIELD_HINTS.items():
            if field in ("subject_code",) or not any(hint.lower() in h for hint in hints):
                continue
            if field == "date" and out["date"]:
                continue
            if field == "deadline" and out["deadline"]:
                continue
            if field in ("org", "code", "subject", "status", "url") and out.get(field):
                continue
            if field == "date":
                out["date"] = value
            elif field == "deadline":
                out["deadline"] = value
            elif field in ("phone", "email"):
                out["phones" if field == "phone" else "emails"] = value
            else:
                out[field] = value
            break
    return out

def parse_ar_meps_table(page_html, base_url, date_order="MDY"):
    items = parse_html_table(page_html, base_url, date_order)
    return items


def parse_generic_links(page_html, base_url):
    items = []
    for href, inner in re.findall(r"<a[^>]+href=[\"']([^\"'#]+)[\"'][^>]*>(.*?)</a>", page_html, re.S | re.I):
        title = strip_tags(inner)
        if len(title) < 15:
            continue
        url = href if href.startswith(("http://", "https://")) else urllib.parse.urljoin(base_url, href)
        items.append({"title": title, "url": url, "snippet": title, "date": "", "deadline": ""})
    uniq, seen_urls = [], set()
    for it in items:
        key = (it["url"], it["title"])
        if key not in seen_urls:
            seen_urls.add(key)
            uniq.append(it)
    return uniq


def parse_rss_items(page_xml):
    items = []
    for block in re.findall(r"<(?:item|entry)[^>]*>(.*?)</(?:item|entry)>", page_xml, re.S | re.I):
        tm = re.search(r"<title[^>]*>(.*?)</title>", block, re.S | re.I)
        title = strip_tags(tm.group(1)) if tm else ""
        lm = re.search(r"<link[^>]*href=[\"']([^\"']+)[\"']", block, re.I) or \
             re.search(r"<link[^>]*>(?:<!\[CDATA\[)?\s*([^>\]]*?)\s*(?:\]\]>)?</link>", block, re.S | re.I)
        dm = re.search(r"<(?:description|summary)[^>]*>(.*?)</(?:description|summary)>", block, re.S | re.I)
        pm = re.search(r"<(?:pubDate|updated|published)[^>]*>(.*?)</(?:pubDate|updated|published)>", block, re.S | re.I)
        items.append({"title": title, "url": lm.group(1).strip() if lm else "",
                      "snippet": strip_tags(dm.group(1))[:600] if dm else title,
                      "date": pm.group(1).strip() if pm else "", "deadline": ""})
    return [i for i in items if i["title"]]


def read_table_file(path, date_order="DMY"):
    """Պահված .xlsx (կամ .csv) աղյուսակը → գրառումներ (Excel-ի «Բեռնել» կոճակի համար)։"""
    items = []
    suffix = path.suffix.lower()
    try:
        if suffix in (".xlsx", ".xlsm"):
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
            for ws in wb.worksheets:
                rows = list(ws.iter_rows(values_only=True))
                if not rows:
                    continue
                header = [str(c).strip() if c is not None else "" for c in rows[0]]
                if sum(1 for c in header if c) < 3:
                    continue
                for r in rows[1:]:
                    cells = ["" if c is None else str(c).strip() for c in r]
                    if not any(cells):
                        continue
                    row = dict(zip(header, cells))
                    context = " | ".join(cells)
                    kind = "bid" if any("մասնակից" in h.lower() or "հվհհ" in h.lower() for h in header) else "tender"
                    items.append({"kind": kind, "row": row, "cells": cells, "text": context,
                                  "title": context[:300], "url": "", "date": "", "deadline": ""})
            wb.close()
        elif suffix in (".csv", ".txt"):
            with path.open(encoding="utf-8-sig", newline="") as fh:
                for row in csv.DictReader(fh):
                    cells = [str(v) for v in row.values()]
                    items.append({"kind": "tender", "row": row, "cells": cells,
                                  "text": " | ".join(cells), "title": " | ".join(cells)[:300],
                                  "url": "", "date": "", "deadline": ""})
    except Exception as exc:
        log.warning("Չհաջողվեց կարդալ %s: %s", path.name, exc)
    return items


# ------------------------------------------------------------------ հիմնական ընթացք
def process_items(items, src, cfg, today, seen, rows, bid_rows, data_dir, page_contacts=None):
    """Մեկ աղբյուրի գտնված գրառումները → ֆայլեր + աղյուսակի տողեր։"""
    name = src.get("name", "կայք")
    date_order = src.get("date_order", "DMY")
    since = parse_date_any(cfg.get("since_date"), "DMY") if cfg.get("since_date") else None
    new_count = 0
    page_dates = []

    for it in items:
        kind = it.get("kind", "tender")
        item_order = it.get("_date_order") or date_order
        row = map_row_fields(it.get("row", {}), date_order) if it.get("row") else {}
        title = (row.get("subject") or it.get("title") or it.get("text", ""))[:300]
        if not title:
            continue
        blob = " ".join([title, it.get("text", ""), it.get("snippet", "")])
        if cfg.get("only_keywords") and not any(k.lower() in blob.lower() for k in cfg["only_keywords"]):
            continue

        # ամսաթիվ. աղյուսակի սյունակից, «տեքստից», կամ փոխանցված դաշտից
        d = (parse_date_any(row.get("date") or it.get("date"), item_order)
             or parse_date_any(extract_fields(blob)["date"], item_order))
        deadline_raw = row.get("deadline") or it.get("deadline") or extract_fields(blob)["deadline"]
        dl = parse_date_any(deadline_raw, item_order)
        if d:
            page_dates.append(d)
        if dl:
            page_dates.append(dl)
        # «սկսած հոկտեմբերից» կանոնը. բաց չենք թողնում այն գնումները, որոնք
        # տպագրվել են ավելի վաղ, բայց հայտերի վերջնաժամկետը դեռ գալիս է։
        if since:
            known = [x for x in (d, dl) if x]
            if known and all(x < since for x in known):
                continue
        item_date_text = human_date(d)

        deadline = row.get("deadline") or it.get("deadline") or extract_fields(blob)["deadline"]
        code = row.get("code") or extract_fields(title + " " + it.get("text", ""))["code"]
        org = row.get("org") or extract_fields(blob)["org"]
        status = row.get("status", "")
        url = row.get("url") or it.get("url") or src.get("url", "")
        name_url = url or title
        item_id = hashlib.md5((name_url + "|" + title[:180]).encode("utf-8")).hexdigest()[:12]
        if item_id in seen:
            continue

        if kind == "bid":
            bid_rows.append({
                "Ամսաթիվ": item_date_text or today,
                "Աղբյուր": name,
                "Մրցույթի ծածկագիր": code,
                "Մրցույթի անվանում": row.get("subject", ""),
                "Մասնակից": it["cells"][1] if len(it["cells"]) > 1 else "",
                "ՀՎՀՀ / Անձնագիր": it["cells"][2] if len(it["cells"]) > 2 else "",
                "Երկիր": it["cells"][3] if len(it["cells"]) > 3 else "",
                "Հայտ": it["cells"][4] if len(it["cells"]) > 4 else "",
            })
            seen[item_id] = today
            continue

        detail_text = ""
        if cfg.get("fetch_detail_pages") and url.startswith("http") and src.get("kind") == "html":
            page = fetch(url)
            if page:
                detail_text = strip_tags(page)
            time.sleep(cfg.get("delay_seconds", 2))
        full_text = " ".join([it.get("snippet", ""), it.get("text", ""), detail_text]).strip()
        fields = extract_fields(full_text + " " + title)
        page_contacts = page_contacts or it.get("_contacts") or {}
        phones = row.get("phones") or fields["phones"] or page_contacts.get("phones", "")
        emails = row.get("emails") or fields["emails"] or page_contacts.get("emails", "")
        if not org:
            org = fields["org"] or page_contacts.get("org", "")

        category = categorize(cfg, title, it.get("text", "") + " " + full_text)
        folder = data_dir / today / category
        folder.mkdir(parents=True, exist_ok=True)
        fname = folder / f"{sanitize(name)}_{sanitize(title)}.txt"
        fname.write_text(
            f"ԱՂԲՅՈՒՐ՝ {name}\nՀՂՈՒՄ՝ {url}\nՀՐԱՊԱՐԱԿՎԱԾ Է՝ {item_date_text or '—'}\n"
            f"ՎԵՐՋՆԱԺԱՄԿԵՏ՝ {deadline or '—'}\nԿԱԶՄԱԿԵՐՊՈՒԹՅՈՒՆ՝ {org or '—'}\n"
            f"ԾԱԾԿԱԳԻՐ՝ {code or '—'}\nՀԵՌԱԽՈՍ՝ {phones or '—'}\nԷԼ. ՓՈՍՏ՝ {emails or '—'}\n"
            f"{'-' * 60}\n{full_text}\n",
            encoding="utf-8",
        )
        rows.append({
            "Ամսաթիվ": item_date_text or today,
            "Աղբյուր": name,
            "Կատեգորիա": category,
            "Ապրանք / Առարկա (նպատակը)": title,
            "Կազմակերպություն": org,
            "Ծածկագիր": code,
            "Վերջնաժամկետ": human_date(dl) or deadline,
            "Կարգավիճակ": status,
            "Հեռախոս": phones,
            "Էլ. փոստ": emails,
            "Հղում": url,
            "Պահպանված ֆայլ": str(fname.relative_to(BASE)),
        })
        seen[item_id] = today
        new_count += 1

    return new_count, page_dates


def write_csv(path, columns, rows, append=False):
    exists = path.exists()
    with path.open("a" if append else "w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=columns, extrasaction="ignore")
        if not append or not exists:
            writer.writeheader()
        writer.writerows(rows)


def process_inbox(cfg, date_order_default="DMY"):
    """inbox/ թղթապանակի .html/.xlsx/.csv ֆայլերը (զննարկիչով պահված էջեր)։"""
    inbox = BASE / cfg.get("inbox_folder", "inbox")
    if not inbox.exists():
        return []
    items = []
    for path in sorted(inbox.iterdir()):
        if path.is_dir() or path.suffix.lower() not in (".html", ".htm", ".xlsx", ".xlsm", ".csv"):
            continue
        log.info("inbox-ից կարդացվում է՝ %s", path.name)
        if path.suffix.lower() in (".html", ".htm"):
            page = path.read_text(encoding="utf-8", errors="replace")
            order = "MDY" if re.search(r"Հրատարակման|Մրցույթի ծածկագիր|ppcm|bid-report", page) else date_order_default
            found = parse_minfin_ads(page, "") or []
            if not found:
                found = parse_html_table(page, "", order)
            if not found:
                found = [{"kind": "tender", **it} for it in parse_generic_links(page, "")]
            contacts = extract_fields(strip_tags(page))
            for it in found:
                it["_date_order"] = order
                it["_contacts"] = contacts
                it.setdefault("url", path.name)
                it["url"] = it.get("url") or path.name
            for it in found:
                it["_file"] = str(path)
            items.extend(found)
        else:
            for it in read_table_file(path, date_order_default):
                it["_file"] = str(path)
                items.append(it)
        processed = inbox / "processed"
        processed.mkdir(exist_ok=True)
        try:
            path.rename(processed / path.name)
        except OSError:
            pass
    return items


def run_once(cfg, demo=False, ignore_seen=False, archive=False, since_override=None):
    today = dt.date.today().isoformat()
    if since_override:
        cfg["since_date"] = since_override
    data_dir = BASE / "data"
    reports_dir = BASE / "reports"
    data_dir.mkdir(exist_ok=True)
    reports_dir.mkdir(exist_ok=True)

    seen = {} if ignore_seen else load_seen()
    rows, bid_rows, per_source, errors = [], [], {}, []

    sources = cfg["sources"]
    if demo:
        sources = [
            {"name": "Օրինակ-ARMEPS-Հայտեր", "url": str(BASE / "samples" / "sample_bid_report.html"),
             "kind": "file", "parse": "table", "date_order": "MDY"},
            {"name": "Օրինակ-MinFin-Հայտարարություն", "url": str(BASE / "samples" / "sample_minfin_ad.html"),
             "kind": "file", "parse": "tender-ad"},
        ]
        log.info("ԴԵՄՈ ռեժիմ. կայքերի փոխարեն կարդացվում են լոկալ օրինակ ֆայլերը։")

    # ---------- inbox (զննարկիչով պահված էջեր/Excel)
    inbox_items = process_inbox(cfg)
    if inbox_items:
        cnt, _ = process_items(inbox_items, {"name": "inbox", "kind": "file"}, cfg, today, seen, rows, bid_rows, data_dir)
        per_source["inbox"] = cnt
        log.info("inbox-ից ավելացվեց %d նոր գրառում", cnt)

    # ---------- կայքեր
    for src in sources:
        name = src.get("name", "կայք")
        kind = src.get("kind", "html")
        mode = src.get("parse", "links")
        date_order = src.get("date_order", "DMY")
        pattern = src.get("page_pattern")
        per_source.setdefault(name, 0)
        log.info("Ստուգվում է՝ %s (%s)", name, src["url"])

        pages = [src["url"]]
        if archive and pattern:
            pages = [pattern.format(n=1)] + [pattern.format(n=n) for n in range(2, cfg.get("max_pages", 25) + 1)]

        for page_url in pages:
            if kind == "file" or page_url.startswith("file://"):
                path = Path(page_url.replace("file://", ""))
                page = path.read_text(encoding="utf-8") if path.exists() else None
            elif kind == "js":
                page = fetch_js(page_url)
                if page is None:
                    page = fetch(page_url)          # վերջին փորձ՝ սովորական հարցումով
            else:
                page = fetch(page_url)
            if page is None:
                if page_url == pages[0]:
                    errors.append(name)
                break

            if mode == "tender-ad":
                items = parse_minfin_ads(page, page_url)
            elif mode == "table":
                items_cells = parse_html_table(page, page_url, date_order)
                items = []
                for it in items_cells:
                    it.setdefault("url", page_url)
                    items.append(it)
            elif kind == "rss":
                items = parse_rss_items(page)
            else:
                items = parse_generic_links(page, page_url)

            log.info("  %s → գտնվեց %d գրառում", page_url, len(items))
            if not items and kind == "js":
                log.warning("  Այս էջի աղյուսակը JavaScript-ով է բեռնվում. տեղադրիր Playwright-ը "
                            "(pip install playwright && playwright install chromium), կամ էջը պահիր "
                            "զննարկիչով (Ctrl+S) և դիր inbox\\ թղթապանակը։")
            page_contacts = extract_fields(strip_tags(page))
            found, page_dates = process_items(items, src, cfg, today, seen, rows, bid_rows,
                                              data_dir, page_contacts)
            per_source[name] += found

            # կանգ առնելը. եթե էջի բոլոր ամսաթվերը հին են, հաջորդ էջերը պետք չեն
            since = parse_date_any(cfg.get("since_date"), "DMY")
            if page_dates and since and max(page_dates) < since:
                log.info("  %s. էջի բոլոր հայտարարությունները ավելի հին են, քան %s-ը, կանգ ենք առնում։",
                         name, cfg.get("since_date"))
                break
            if kind == "file":
                break
            time.sleep(cfg.get("delay_seconds", 2))
            if not archive and len(pages) == 1:
                break

    SEEN_PATH.write_text(json.dumps(seen, ensure_ascii=False, indent=0), encoding="utf-8")

    # ---------- հաշվետվություններ
    if rows:
        write_csv(reports_dir / f"Հաշվետվություն-{today}.csv", TENDER_COLUMNS, rows)
        write_csv(reports_dir / "MASTER.csv", TENDER_COLUMNS, rows, append=True)
    if bid_rows:
        write_csv(reports_dir / f"Հայտեր-{today}.csv", BID_COLUMNS, bid_rows)
        write_csv(reports_dir / "MASTER-Հայտեր.csv", BID_COLUMNS, bid_rows, append=True)

    summary = [f"ԱՄՓՓԱԳԻՐ {today}", "=" * 44,
               f"Նոր հայտարարություններ՝ {len(rows)}   |   Հայտեր՝ {len(bid_rows)}"]
    for name, cnt in per_source.items():
        summary.append(f"  • {name}: {cnt} նոր")
    if cfg.get("since_date"):
        summary.append(f"  (հաշվի են առնվում {cfg['since_date']}-ից սկսած հայտարարությունները)")
    if errors:
        summary.append("Չբացված կայքեր՝ " + ", ".join(errors))
    for row in rows:
        summary.append(f"  - [{row['Կատեգորիա']}] {row['Ապրանք / Առարկա (նպատակը)'][:64]}"
                       f" | հեռ.՝ {row['Հեռախոս'] or '—'} | էլ.փ.՝ {row['Էլ. փոստ'] or '—'}")
    text = "\n".join(summary)
    (reports_dir / f"Ամփոփագիր-{today}.txt").write_text(text, encoding="utf-8")
    log.info("\n%s", text)
    return rows, bid_rows


def loop_mode(cfg):
    target = cfg.get("run_time", "08:00")
    log.info("Մշտական ռեժիմ. ամեն օր %s-ին կստուգի կայքերը։ Փակելու համար՝ Ctrl+C։", target)
    while True:
        now = dt.datetime.now()
        hh, mm = map(int, target.split(":"))
        nxt = now.replace(hour=hh, minute=mm, second=0, microsecond=0)
        if nxt <= now:
            nxt += dt.timedelta(days=1)
        log.info("Հաջորդ գործարկումը՝ %s", nxt.strftime("%Y-%m-%d %H:%M"))
        time.sleep((nxt - now).total_seconds())
        try:
            run_once(cfg)
        except Exception as exc:
            log.exception("Սխալ գործարկման ժամանակ. %s", exc)


def main():
    parser = argparse.ArgumentParser(description="Տենդերային մոնիտոր")
    parser.add_argument("--demo", action="store_true", help="փորձնական ռեժիմ լոկալ օրինակներով")
    parser.add_argument("--archive", action="store_true", help="շրջանցել բոլոր էջերը՝ since_date-ից սկսած")
    parser.add_argument("--since", help="սկսած ամսաթիվ, օր.՝ 2026-10-01")
    parser.add_argument("--loop", action="store_true", help="ամենօրյա ավտոմատ ռեժիմ")
    parser.add_argument("--reset", action="store_true", help="մոռանալ հին բեռնումները")
    args = parser.parse_args()

    setup_logging()
    cfg = load_config()
    if args.reset and SEEN_PATH.exists():
        SEEN_PATH.unlink()
        log.info("Հիշողությունը մաքրված է։")
    if args.loop:
        loop_mode(cfg)
    else:
        rows, bids = run_once(cfg, demo=args.demo, ignore_seen=args.reset,
                              archive=args.archive, since_override=args.since)
        log.info("Ավարտվեց։ Հայտարարություններ՝ %d, հայտեր՝ %d։ Տեսեք reports/ և data/ թղթապանակները։",
                 len(rows), len(bids))


if __name__ == "__main__":
    main()
